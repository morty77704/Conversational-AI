from __future__ import annotations

import io
from pathlib import Path

import pymupdf
import pytesseract
from bs4 import BeautifulSoup
from docx import Document as DocxDocument
from langchain_core.documents import Document
from openpyxl import load_workbook
from PIL import Image
from pypdf import PdfReader

from common.config import get_settings

TEXT_EXTENSIONS = {".md", ".txt"}
HTML_EXTENSIONS = {".html", ".htm"}
IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".bmp", ".tif", ".tiff"}
SUPPORTED_EXTENSIONS = (
    TEXT_EXTENSIONS
    | HTML_EXTENSIONS
    | IMAGE_EXTENSIONS
    | {".pdf", ".docx", ".xlsx"}
)
MIN_PDF_TEXT_LENGTH = 40


def configure_tesseract() -> None:
    """设置本地 Tesseract 路径，并在真正需要 OCR 时验证文件。"""
    tesseract_cmd = get_settings().tesseract_cmd

    if not tesseract_cmd.exists():
        raise FileNotFoundError(f"Tesseract 不存在：{tesseract_cmd}")

    pytesseract.pytesseract.tesseract_cmd = str(tesseract_cmd)


def normalize_text(text: str) -> str:
    lines = [line.rstrip() for line in text.replace("\r\n", "\n").split("\n")]
    normalized_lines: list[str] = []
    previous_blank = False

    for line in lines:
        is_blank = not line.strip()

        if is_blank and previous_blank:
            continue

        normalized_lines.append(line)
        previous_blank = is_blank

    return "\n".join(normalized_lines).strip()


def load_text_file(path: Path) -> str:
    return normalize_text(path.read_text(encoding="utf-8-sig"))


def load_html_file(path: Path) -> str:
    html = path.read_text(encoding="utf-8-sig")
    soup = BeautifulSoup(html, "html.parser")

    for tag in soup(["script", "style", "noscript"]):
        tag.decompose()

    return normalize_text(soup.get_text("\n"))


def load_docx_file(path: Path) -> str:
    docx = DocxDocument(path)
    blocks = [paragraph.text for paragraph in docx.paragraphs if paragraph.text.strip()]

    for table in docx.tables:
        for row in table.rows:
            values = [cell.text.strip() for cell in row.cells]
            blocks.append(" | ".join(values))

    return normalize_text("\n\n".join(blocks))


def load_xlsx_file(path: Path) -> str:
    workbook = load_workbook(path, read_only=True, data_only=False)
    blocks: list[str] = []

    try:
        for sheet in workbook.worksheets:
            blocks.append(f"# 工作表：{sheet.title}")

            for row_index, row in enumerate(sheet.iter_rows(values_only=True), start=1):
                values = ["" if value is None else str(value).strip() for value in row]

                if any(values):
                    blocks.append(f"第 {row_index} 行 | " + " | ".join(values))
    finally:
        workbook.close()

    return normalize_text("\n".join(blocks))


def ocr_image(image: Image.Image) -> str:
    configure_tesseract()
    settings = get_settings()
    return normalize_text(
        pytesseract.image_to_string(
            image,
            lang=settings.ocr_languages,
        )
    )


def load_image_file(path: Path) -> str:
    with Image.open(path) as image:
        return ocr_image(image.convert("RGB"))


def ocr_pdf_page(path: Path, page_index: int) -> str:
    configure_tesseract()

    with pymupdf.open(path) as pdf:
        page = pdf.load_page(page_index)
        pixmap = page.get_pixmap(dpi=200, alpha=False)
        image = Image.open(io.BytesIO(pixmap.tobytes("png"))).convert("RGB")
        return ocr_image(image)


def load_pdf_file(path: Path) -> str:
    reader = PdfReader(str(path))
    pages: list[str] = []

    for page_index, page in enumerate(reader.pages):
        extracted_text = normalize_text(page.extract_text() or "")

        # 仅在页面几乎没有可提取文字时启用 OCR，避免无意义地增加耗时。
        if len(extracted_text.replace(" ", "")) < MIN_PDF_TEXT_LENGTH:
            extracted_text = ocr_pdf_page(path, page_index)

        pages.append(f"# 第 {page_index + 1} 页\n\n{extracted_text}")

    return normalize_text("\n\n".join(pages))


def load_document(path: Path, metadata: dict[str, object] | None = None) -> Document:
    resolved_path = path.resolve()

    if not resolved_path.exists():
        raise FileNotFoundError(f"知识库源文件不存在：{resolved_path}")

    extension = resolved_path.suffix.lower()

    if extension not in SUPPORTED_EXTENSIONS:
        raise ValueError(f"暂不支持的知识库文件类型：{extension}")

    if extension in TEXT_EXTENSIONS:
        content = load_text_file(resolved_path)
    elif extension in HTML_EXTENSIONS:
        content = load_html_file(resolved_path)
    elif extension == ".pdf":
        content = load_pdf_file(resolved_path)
    elif extension == ".docx":
        content = load_docx_file(resolved_path)
    elif extension == ".xlsx":
        content = load_xlsx_file(resolved_path)
    else:
        content = load_image_file(resolved_path)

    if not content:
        raise ValueError(f"源文件解析后没有文字：{resolved_path}")

    document_metadata = dict(metadata or {})
    document_metadata.setdefault("file_path", str(resolved_path))
    document_metadata.setdefault("format", extension.removeprefix("."))

    return Document(
        page_content=content,
        metadata=document_metadata,
    )
